from abc import ABC
from datetime import timedelta
from typing import Iterable
from typing import Optional
from typing import cast

# from typing import NoReturn
# from typing import Sequence

from lxml.etree import _Element
from openpyxl.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from package.utilities import make_id_cells
from package.utilities import CellT
from package.utilities import Excel
from package.utilities import BOLD
from package.utilities import counters


class SectionParent:
    """Some sections are grouped together under subheadings in the
    table of contents. This is usefull for calculating the total
    duration and 'after appointed time' for those subheadings."""

    def __init__(self, title: str):
        self.title = title

        self.total_duration = timedelta()
        self.total_aat = timedelta()


class TableSection(ABC):
    def __init__(
        self,
        title: str,
        excel_sheet_title: Optional[str] = None,  # not sure about this one
        parent: Optional[SectionParent] = None,
    ):

        # just for typing purposes.
        # really not sure about this one
        self.parent = parent

        self.title = title
        self.rows = 0
        self.cells: list[_Element] = []

    def __len__(self):
        return len(self.cells)

    def add_row(self, cells_items: Iterable[CellT]):
        self.rows += 1
        cells = make_id_cells(cells_items)
        self.cells.extend(cells)
        # if counters.tables_abc_add_row < 3:
        #     print(f"{cells=}")
        #     print(f"{self.cells=}")
        #     counters.tables_abc_add_row += 1

    # def add_to(self, table: Analysis_Table):
    #     # ID XML stuff
    #     table.add_table_sub_head(self.title)
    #     table.extend(self.cells)
    #     table.increment_rows(increment_by=self.rows)


class AnalysisTableSection(TableSection):
    def __init__(
        self,
        title: str,
        excel_sheet_title: Optional[str] = None,
        parent: Optional[SectionParent] = None,
    ):
        super().__init__(title, excel_sheet_title, parent)
        self.parent = parent
        self.duration = timedelta(seconds=0)
        self.after_appointed_time: timedelta = timedelta(seconds=0)

        # also create an excel sheet
        self.excel_sheet: Optional[Worksheet] = None
        if Excel.out_wb is not None:
            self.excel_sheet = cast(
                Worksheet, Excel.out_wb.create_sheet(excel_sheet_title)
            )

    # ignoring The Liskov Substitution Principle here
    def add_row(
        self,
        cells_items: Iterable[CellT],
        duration: timedelta,
        aat: Optional[timedelta] = None,
    ):

        super().add_row(cells_items)

        self.duration += duration

        if aat:
            self.after_appointed_time += aat

        e_row: list[CellT] = []
        for cell in cells_items:
            if isinstance(cell, str):
                cell = cell.replace("\t", " ")
                e_row.append(cell)
            else:
                e_row.append(cell)
        # print(e_row)
        if self.excel_sheet:
            self.excel_sheet.append(e_row)

    def add_to_excel(self) -> None:

        # excel stuff
        if self.excel_sheet:
            sess_tot_cell = Cell(self.excel_sheet, value="Sessional Total")
            sess_tot_cell.font = BOLD
            tot_dur_cell = Cell(self.excel_sheet, value=self.duration)  # type: ignore
            tot_dur_cell.font = BOLD

            totals_row = [None, sess_tot_cell, tot_dur_cell]

            self.excel_sheet.insert_rows(1, 2)
            self.excel_sheet["A1"] = self.title.replace("\t", " ")
            # make first row bold
            self.excel_sheet["A1"].font = BOLD  # type: ignore
            for cell in self.excel_sheet[2]:  # get second row
                cell.font = BOLD  # make second row bold
            self.excel_sheet["A2"] = "Date"
            self.excel_sheet["B2"] = "Content"
            self.excel_sheet["C2"] = "Duration"

            self.excel_sheet.append(totals_row)

            # tidy up the col widths of the first two columns. Otherwise
            # it's too narrow and you have to change it every time you
            # open the excel
            self.excel_sheet.column_dimensions["A"].width = 20  # type: ignore
            self.excel_sheet.column_dimensions["B"].width = 30  # type: ignore


class WH_DiaryDay_TableSection(TableSection):
    def __init__(self, title: str):
        super().__init__(title)
        self.duration = timedelta(seconds=0)

    def add_row(self, cells: Iterable[CellT], duration: timedelta):
        super().add_row(cells)
        self.duration += duration


class CH_DiaryDay_TableSection(WH_DiaryDay_TableSection):
    def __init__(self, title: str):
        super().__init__(title)
        self.after_appointed_time = timedelta(seconds=0)

    def add_row(
        self,
        cells: Iterable[CellT],
        duration: timedelta,
        aat: timedelta,
    ):
        super().add_row(cells, duration)
        self.after_appointed_time += aat

        # if counters.diary_add_row < 3:
        #     print(f"diary add row:\n  {cells=}\n  {duration=}\n  {aat=}")
        #     counters.diary_add_row += 1

    def __str__(self):
        return f"{self.__class__}({self.title})\n  {self.cells=}"

    # def add_to(
    #     self,  # type: ignore
    #     table: CH_Diary_Table,
    #     session_duration: timedelta,
    #     session_aat: timedelta,
    # ):

    #     table.add_table_sub_head(self.title)
    #     table.extend(self.cells)
    #     table.increment_rows(increment_by=self.rows)
    #     table.add_total_duration(
    #         self.duration, self.after_appointed_time, session_duration, session_aat
    #     )


# class AnalysisTableSection(AnalysisTableSection):
#     def __init__(
#         self, title: str, excel_sheet_title: str, parent: Optional[SectionParent]
#     ):
#         super().__init__(title, excel_sheet_title, parent)
#         self.after_appointed_time = timedelta(seconds=0)

#     def add_row(
#         self,  # type: ignore
#         cells_items: Iterable[CellT],
#         duration: timedelta,
#         aat: timedelta,
#     ):

#         super().add_row(cells_items, duration=duration)
#         self.after_appointed_time += aat
