from datetime import timedelta
from typing import Iterable, Optional, cast, NoReturn

from openpyxl.cell import Cell
from openpyxl.styles import Font
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from lxml import etree
from lxml.etree import SubElement
from lxml.etree import _Element

import package.utilities as utils
from package.utilities import AID, AID5
from package.utilities import make_id_cells
from package.utilities import format_timedelta
from package.utilities import CellT

BOLD = Font(bold=True)

# exporting to excel is optional
class Excel:
    # the workbook class to output to
    out_wb: Optional[Workbook] = None


class WH_Table(etree.ElementBase):
    """class is based on etree.ElementBase
    name of the xml element will default to the name of the class
    class can be instantiated despite what type checkers may think"""

    def increment_rows(self, increment_by: int = 1):
        trows = self.get(AID + "trows", default=None)
        if trows:
            self.set(AID + "trows", str(int(trows) + increment_by))
        else:
            self.set(AID + "trows", "1")

    def add_total_duration(self, total_duration: timedelta):
        self.increment_rows()
        total_cell = utils.Body_line_below_right_align()
        total_cell.text = "Total:"
        time_cell = utils.Body_lines()
        time_cell.text = format_timedelta(total_duration)
        self.extend(make_id_cells([None]) + [total_cell, time_cell])

    def add_table_sub_head(self, heading_text: str, subsubhead: bool = False):
        self.increment_rows()
        cellstyle = "SubHeading"

        if subsubhead:
            cellstyle = "SubSubHeading"

        sub_head = SubElement(
            self,
            "Cell",
            attrib={
                AID + "table": "cell",
                AID + "ccols": "3",
                AID5 + "cellstyle": cellstyle,
            },
        )
        sub_head.text = heading_text


class _TableSection:
    def __init__(self, title: str):

        # just for typing purposes.
        # really not sure about this one
        self.parent = None

        self.title = title
        self.rows = 0
        self.cells: list[_Element] = []

    def __len__(self):
        return len(self.cells)

    def add_row(self, cells_items: Iterable[CellT]):
        self.rows += 1
        cells = make_id_cells(cells_items)
        self.cells.extend(cells)

    def add_to(self, table: WH_Table):
        # ID XML stuff
        table.add_table_sub_head(self.title)
        table.extend(self.cells)
        table.increment_rows(increment_by=self.rows)


class SudoTableSection(_TableSection):
    """These are just to be used in the table of contents"""

    def __init__(self, title: str):
        self.title = title

        self.total_duration = timedelta()
        self.total_aat = timedelta()

    def __len__(self):
        return 0

    def add_row(self, cells_items: Iterable[CellT]) -> NoReturn:
        raise NotImplementedError

    def add_to(self, table: WH_Table) -> NoReturn:
        raise NotImplementedError


class WH_AnalysisTableSection(_TableSection):

    # for 'Part' totals on the contents page
    part_dur = timedelta(seconds=0)

    def __init__(
        self, title: str, excel_sheet_title: str, parent: Optional[SudoTableSection]
    ):
        super().__init__(title)
        self.parent = parent
        self.duration = timedelta(seconds=0)
        # also create an excel sheet
        self.excel_sheet: Optional[Worksheet] = None
        if Excel.out_wb:
            self.excel_sheet = cast(
                Worksheet, Excel.out_wb.create_sheet(excel_sheet_title)
            )

    # ignoring The Liskov Substitution Principle here
    def add_row(
        self, cells_items: Iterable[CellT], duration: timedelta  # type: ignore
    ):

        super().add_row(cells_items)

        if duration is not None:
            self.duration += duration

        e_row: list[CellT] = []
        for cell in cells_items:
            if isinstance(cell, str):
                cell = cell.replace("\t", " ")
                e_row.append(cell)
            else:
                e_row.append(cell)
        # print(e_row)
        if self.excel_sheet:
            self.excel_sheet.append(e_row)

    def _add_to_excel(self, totals_row: list[None | Cell]) -> None:
        if self.excel_sheet:
            self.excel_sheet.insert_rows(1, 2)
            self.excel_sheet["A1"] = self.title.replace("\t", " ")
            # make first row bold
            self.excel_sheet["A1"].font = BOLD  # type: ignore
            for cell in self.excel_sheet[2]:  # get second row
                cell.font = BOLD  # make second row bold
            self.excel_sheet["A2"] = "Date"
            self.excel_sheet["B2"] = "Content"
            self.excel_sheet["C2"] = "Duration"

            self.excel_sheet.append(totals_row)

            # tidy up the col widths of the first two columns.
            # Otherwise it's too narrow and you have to change it every time you open the excel
            self.excel_sheet.column_dimensions["A"].width = 20
            self.excel_sheet.column_dimensions["B"].width = 30

    def add_to(self, table: WH_Table):
        # super().add_to(table)
        # table.add_total_duration(self.duration)

        if self.parent is not None:
            # if there is a parent then this section of the table should have
            # a subsubsection heading rather than a subsection heading
            table.add_table_sub_head(self.title, subsubhead=True)
        else:
            table.add_table_sub_head(self.title)

        table.extend(self.cells)
        table.increment_rows(increment_by=self.rows)
        table.add_total_duration(self.duration)

        if self.parent is not None:
            self.parent.total_duration += self.duration

        # excel stuff
        if self.excel_sheet:
            sess_tot_cell = Cell(self.excel_sheet, value="Sessional Total")
            sess_tot_cell.font = BOLD
            tot_dur_cell = Cell(self.excel_sheet, value=self.duration)  # type: ignore
            tot_dur_cell.font = BOLD

            totals_row = [None, sess_tot_cell, tot_dur_cell]

            self._add_to_excel(totals_row)


class CH_Table(WH_Table):
    def add_total_duration(
        self, total_duration: timedelta, aat_total: timedelta  # type: ignore
    ) -> None:

        self.increment_rows()
        total_cell = utils.Body_line_below_right_align()
        total_cell.text = "Total:"

        time_cell = utils.Body_lines()
        time_cell.text = format_timedelta(total_duration)

        time_2_cell = utils.Body_lines()
        time_2_cell.text = format_timedelta(aat_total)

        self.extend(make_id_cells([None]) + [total_cell, time_cell, time_2_cell])

    def add_table_sub_head(self, heading_text: str, subsubhead: bool = False):
        self.increment_rows()
        cellstyle = "SubHeading"

        if subsubhead:
            cellstyle = "SubSubHeading"

        sub_head = SubElement(
            self,
            "Cell",
            attrib={
                AID + "table": "cell",
                AID + "ccols": "4",
                AID5 + "cellstyle": cellstyle,
            },
        )
        sub_head.text = heading_text


class CH_AnalysisTableSection(WH_AnalysisTableSection):

    # create some text for the totals on the contents page
    # contents_text = ''
    part_aat = timedelta(seconds=0)
    table_num_aat = {}

    def __init__(
        self, title: str, excel_sheet_title: str, parent: Optional[SudoTableSection]
    ):
        super().__init__(title, excel_sheet_title, parent)
        self.after_appointed_time = timedelta(seconds=0)

    def add_row(
        self,  # type: ignore
        cells_items: Iterable[CellT],
        duration: timedelta,
        aat: timedelta,
    ):

        super().add_row(cells_items, duration=duration)
        self.after_appointed_time += aat

    def add_to(self, table: CH_Table):  # type: ignore
        if self.parent is not None:
            # if there is a parent then this section of the
            # table should have a subsubsection heading rather
            # than a subsection heading
            table.add_table_sub_head(self.title, subsubhead=True)
        else:
            table.add_table_sub_head(self.title)
        table.extend(self.cells)
        table.increment_rows(increment_by=self.rows)
        table.add_total_duration(self.duration, self.after_appointed_time)

        CH_AnalysisTableSection.part_dur += self.duration
        CH_AnalysisTableSection.part_aat += self.after_appointed_time

        # some sections have parents referenced in the table of contents
        # these parents also need to have the durations calculated
        if self.parent is not None:
            self.parent.total_duration += self.duration
            self.parent.total_aat += self.after_appointed_time

        # excel stuff
        if self.excel_sheet:
            sess_tot_cell = Cell(self.excel_sheet, value="Sessional Total")
            sess_tot_cell.font = BOLD
            tot_dur_cell = Cell(self.excel_sheet, value=self.duration)  # type: ignore
            tot_dur_cell.font = BOLD
            tot_aat_cell = Cell(self.excel_sheet, value=self.after_appointed_time)  # type: ignore
            tot_aat_cell.font = BOLD

            totals_row = [None, sess_tot_cell, tot_dur_cell, tot_aat_cell]

            self._add_to_excel(totals_row)

    def _add_to_excel(self, totals_row: list[Optional[Cell]]):
        if self.excel_sheet:
            super()._add_to_excel(totals_row)
            # the chamber is different from WH as it includes an extra col
            self.excel_sheet["D2"] = "After appointed time"  # extra col head


class WH_Diary_Table(WH_Table):

    # ignoring The Liskov Substitution Principle here as I can't think
    # how else I would do this.
    def add_total_duration(
        self,  # type: ignore
        daily_total_duration: timedelta,
        session_total_duration: timedelta,
    ):

        self.increment_rows()
        total_cell = utils.Right_align_cell()
        total_cell.set(AID + "ccols", "2")  # span 2 cols
        total_cell.text = "Daily Totals:"

        time_cell = utils.Body_line_above()
        time_cell.text = format_timedelta(daily_total_duration)

        self.extend([total_cell, time_cell])  # type: ignore

        self.increment_rows()
        total_cell = utils.Body_line_below_right_align()
        total_cell.set(AID + "ccols", "2")  # span 2 cols
        total_cell.text = "Totals for Session:"
        # print(etree.tostring(total_cell))

        time_cell = utils.Body_line_below()
        time_cell.text = format_timedelta(session_total_duration)
        # print(etree.tostring(time_cell), '\n')

        self.extend([total_cell, time_cell])

    def add_table_sub_head(self, heading_text: str):  # type: ignore
        self.increment_rows()
        sub_head = SubElement(
            self,
            "Cell",
            attrib={
                AID + "table": "cell",
                AID + "ccols": "3",
                AID5 + "cellstyle": "SubHeading No Toc",
            },
        )
        sub_head.text = heading_text


class WH_DiaryDay_TableSection(_TableSection):
    def __init__(self, title: str):
        super().__init__(title)
        self.duration = timedelta(seconds=0)

    def add_row(self, cells: Iterable[CellT], duration: timedelta):  # type: ignore
        super().add_row(cells)
        self.duration += duration

    def add_to(
        self, table: WH_Diary_Table, session_total_time: timedelta  # type: ignore
    ):
        table.add_table_sub_head(self.title)
        table.extend(self.cells)
        table.increment_rows(increment_by=self.rows)
        table.add_total_duration(self.duration, session_total_time)


class CH_Diary_Table(WH_Table):
    def add_total_duration(
        self,
        daily_total_duration: timedelta,
        daily_aat_total: timedelta,
        session_total_duration: timedelta,
        session_aat_total: timedelta,
    ) -> None:

        self.increment_rows()
        total_cell = utils.Right_align_cell()
        # total_cell.set(AID + 'ccols', '2')  # span 2 cols
        total_cell.text = "Daily Totals:"

        time_cell = utils.Body_line_above()
        time_cell.text = format_timedelta(daily_total_duration)

        time_2_cell = utils.Body_line_above()
        time_2_cell.text = format_timedelta(daily_aat_total)

        self.extend(make_id_cells([None]) + [total_cell, time_cell, time_2_cell])

        self.increment_rows()
        total_cell = utils.Body_line_below_right_align()
        # total_cell.set(AID + 'ccols', '2')  # span 2 cols
        total_cell.text = "Totals for Session:"

        time_cell = utils.Body_line_below()
        time_cell.text = format_timedelta(session_total_duration)
        time_2_cell = utils.Body_line_below()
        time_2_cell.text = format_timedelta(session_aat_total)

        cells = make_id_cells([None], attrib={AID5 + "cellstyle": "BodyLineBelow"}) + [
            total_cell,
            time_cell,
            time_2_cell,
        ]
        self.extend(cells)

    def add_table_sub_head(self, heading_text: str):
        self.increment_rows()
        SubElement(
            self,
            "Cell",
            attrib={
                AID + "table": "cell",
                AID + "ccols": "4",
                AID5 + "cellstyle": "SubHeading No Toc",
            },
        ).text = heading_text


class CH_DiaryDay_TableSection(WH_DiaryDay_TableSection):
    def __init__(self, title: str):
        super().__init__(title)
        self.after_appointed_time = timedelta(seconds=0)

    def add_row(
        self,  # type: ignore
        cells: Iterable[CellT],
        duration: timedelta,
        aat: timedelta,
    ):
        super().add_row(cells, duration)
        self.after_appointed_time += aat

    def add_to(
        self,  # type: ignore
        table: CH_Diary_Table,
        session_duration: timedelta,
        session_aat: timedelta,
    ):

        table.add_table_sub_head(self.title)
        table.extend(self.cells)
        table.increment_rows(increment_by=self.rows)
        table.add_total_duration(
            self.duration, self.after_appointed_time, session_duration, session_aat
        )


class Contents_Table(WH_Table):
    # def add_row(self, cells):
    #     self.increment_rows()
    #     self.extend(deepcopy(cells))

    def add_row(self, cells_items: Iterable[CellT]):
        self.increment_rows()
        cells = make_id_cells(cells_items)
        self.extend(cells)
